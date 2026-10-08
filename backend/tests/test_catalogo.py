"""Catálogo de cámaras, sedes y resumen del Inicio."""

import datetime as dt
import io

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import config
from app.main import app
from app.services import catalogo, localidades

from .municipios_de_prueba import escribir

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAPEO = {"col_id": "Cámara", "col_lat": "Lat", "col_lon": "Lon", "col_label": "Dirección", "col_node": "Nodo"}


@pytest.fixture
def client() -> TestClient:
    return TestClient(app)


@pytest.fixture
def municipios() -> None:
    escribir(config.MUNICIPIOS_PATH)


def _planilla(filas: list[dict]) -> bytes:
    buffer = io.BytesIO()
    pd.DataFrame(filas).to_excel(buffer, index=False, engine="openpyxl")
    return buffer.getvalue()


def _fila(camara: str, lat: float, lon: float, direccion: str | None = None, nodo: str | None = None) -> dict:
    return {"Cámara": camara, "Lat": lat, "Lon": lon, "Dirección": direccion, "Nodo": nodo}


BASE = [
    _fila("C-1", -24.75, -65.45, "Av. Belgrano 100", "N1"),
    _fila("C-2", -24.80, -65.40, "Enclave 5"),
    _fila("C-3", -24.95, -65.40, "Ruta 68 km 3"),
]


def _importar(
    client: TestClient, filas: list[dict] | bytes, nombre: str = "camaras.xlsx", **mapeo
) -> dict:
    contenido = filas if isinstance(filas, bytes) else _planilla(filas)
    response = client.post(
        "/catalogo/importar/",
        files={"file": (nombre, contenido, XLSX_MIME)},
        data={**MAPEO, **mapeo},
    )
    assert response.status_code == 200, response.text
    return response.json()


def _camaras(client: TestClient) -> dict[str, dict]:
    response = client.get("/catalogo/camaras/")
    assert response.status_code == 200, response.text
    return {camara["id"]: camara for camara in response.json()}


# --------------------------------------------------------------------------
# Importar
# --------------------------------------------------------------------------


def test_importar_agrega_las_camaras_con_su_localidad(client, municipios) -> None:
    resultado = _importar(client, BASE + [_fila("C-mala", 0, 0)])

    assert resultado["added"] == 3
    assert resultado["valid"] == 3 and resultado["rows"] == 4
    assert [fila["row"] for fila in resultado["discarded"]] == [4]
    assert resultado["camera_count"] == 3
    assert resultado["municipalities_loaded"] is True

    camaras = _camaras(client)
    assert {c: camaras[c]["locality"] for c in camaras} == {
        "C-1": "Salta", "C-2": "Enclave", "C-3": "Cerrillos",
    }
    assert camaras["C-1"]["label"] == "Av. Belgrano 100"
    assert camaras["C-1"]["node"] == "N1"
    assert camaras["C-1"]["visits"] == 0 and camaras["C-1"]["last_status"] is None


def test_importar_combina_por_id_sin_borrar(client, municipios) -> None:
    _importar(client, BASE)
    resultado = _importar(
        client,
        [
            _fila("C-1", -24.76, -65.45),  # se movió; sin dirección: no la borra
            _fila("C-2", -24.80, -65.40, "Enclave 5"),  # igual
            _fila("C-4", -23.15, -64.35, "Orán"),  # nueva
            _fila("C-4", -22.50, -63.80, "Orán (corregida)"),  # repetida: vale la última
        ],
        nombre="otra.xlsx",
    )

    assert (resultado["added"], resultado["updated"], resultado["unchanged"]) == (1, 1, 1)
    assert resultado["duplicates"] == 1
    assert resultado["camera_count"] == 4  # C-3 no vino y sigue estando

    camaras = _camaras(client)
    assert camaras["C-1"]["lat"] == -24.76
    assert camaras["C-1"]["label"] == "Av. Belgrano 100"
    assert camaras["C-4"]["label"] == "Orán (corregida)"
    assert camaras["C-4"]["locality"] == "Islas"


def test_el_mismo_archivo_dos_veces_avisa_y_no_cambia_nada(client) -> None:
    # Los mismos bytes: el .xlsx lleva la hora en que se generó.
    contenido = _planilla(BASE)
    primero = _importar(client, contenido)
    segundo = _importar(client, contenido)

    assert primero["previously_loaded_at"] is None
    assert segundo["previously_loaded_at"] is not None
    assert (segundo["added"], segundo["updated"], segundo["unchanged"]) == (0, 0, 3)
    historial = client.get("/catalogo/importaciones/").json()
    assert [carga["id"] for carga in historial] == [segundo["import_id"], primero["import_id"]]
    assert historial[0]["rows"] == 3


def test_sin_limites_queda_sin_calcular_y_se_recalcula_al_tenerlos(client) -> None:
    resultado = _importar(client, BASE)
    assert resultado["municipalities_loaded"] is False
    assert {c["locality"] for c in _camaras(client).values()} == {localidades.SIN_CALCULAR}

    escribir(config.MUNICIPIOS_PATH)
    # Cualquier importación completa las que quedaron sin calcular: las que no
    # vienen (C-3) y las que vienen igual (C-2) o cambian sólo el texto (C-1).
    _importar(client, [_fila("C-9", -24.75, -65.45), BASE[1], {**BASE[0], "Dirección": "Otra"}])
    camaras = _camaras(client)
    assert {c: camaras[c]["locality"] for c in ("C-1", "C-2", "C-3", "C-9")} == {
        "C-1": "Salta", "C-2": "Enclave", "C-3": "Cerrillos", "C-9": "Salta",
    }


def test_importar_calcula_la_localidad_solo_de_lo_nuevo_o_movido(client, municipios, monkeypatch) -> None:
    """Con miles de cámaras, ubicar todas en cada importación tardaba segundos
    dentro de la transacción. Sólo cambian de municipio las nuevas y las que
    se movieron (y las que quedaron sin calcular)."""
    _importar(client, BASE)
    calculadas: list[tuple[float, float]] = []
    original = localidades.Limites.localidad_de

    def contando(self, lat, lon):
        calculadas.append((lat, lon))
        return original(self, lat, lon)

    monkeypatch.setattr(localidades.Limites, "localidad_de", contando)

    _importar(
        client,
        [
            _fila("C-1", -24.76, -65.45),  # se movió
            _fila("C-2", -24.80, -65.40, "Enclave 5"),  # igual
            _fila("C-3", -24.95, -65.40, "Otra dirección"),  # cambió sólo el texto
            _fila("C-4", -22.50, -63.80),  # nueva
        ],
        nombre="segunda.xlsx",
    )

    assert sorted(calculadas) == sorted([(-24.76, -65.45), (-22.50, -63.80)])
    assert _camaras(client)["C-4"]["locality"] == "Islas"


def test_la_localidad_corregida_a_mano_no_la_pisa_una_importacion(client, municipios) -> None:
    _importar(client, BASE)
    response = client.patch("/catalogo/camaras/C-1", json={"locality": "  Villa San Lorenzo "})
    assert response.status_code == 200
    assert response.json()["locality"] == "Villa San Lorenzo"
    assert response.json()["locality_manual"] is True

    _importar(client, [_fila("C-1", -24.74, -65.45)], nombre="mueve.xlsx")
    assert _camaras(client)["C-1"]["locality"] == "Villa San Lorenzo"

    # null vuelve al cálculo.
    response = client.patch("/catalogo/camaras/C-1", json={"locality": None})
    assert response.json()["locality"] == "Salta"
    assert response.json()["locality_manual"] is False


def test_dar_de_baja_una_camara(client) -> None:
    _importar(client, BASE)
    assert client.delete("/catalogo/camaras/C-2").status_code == 204
    assert set(_camaras(client)) == {"C-1", "C-3"}
    assert client.delete("/catalogo/camaras/C-2").status_code == 404
    assert client.patch("/catalogo/camaras/C-2", json={"locality": "X"}).status_code == 404


def test_ids_con_barra_funcionan_en_la_ruta(client) -> None:
    _importar(client, [_fila("OT/15", -24.75, -65.45)])
    assert client.patch("/catalogo/camaras/OT/15", json={"locality": "Salta"}).status_code == 200
    assert client.delete("/catalogo/camaras/OT/15").status_code == 204


def test_importar_sin_coordenadas_validas_es_422(client) -> None:
    response = client.post(
        "/catalogo/importar/",
        files={"file": ("x.xlsx", _planilla([_fila("C-1", 0, 0)]), XLSX_MIME)},
        data=MAPEO,
    )
    assert response.status_code == 422
    assert "Ninguna de las 1 filas" in response.json()["detail"]


# --------------------------------------------------------------------------
# Permisos
# --------------------------------------------------------------------------


def test_un_operador_consulta_pero_no_modifica(client, como_operador) -> None:
    assert client.get("/catalogo/camaras/").status_code == 200
    assert client.get("/catalogo/sedes/").status_code == 200
    assert client.get("/catalogo/resumen").status_code == 200

    importar = client.post(
        "/catalogo/importar/",
        files={"file": ("c.xlsx", _planilla(BASE), XLSX_MIME)},
        data=MAPEO,
    )
    sede = {"name": "Base", "lat": -24.78, "lon": -65.41}
    for response in (
        importar,
        client.patch("/catalogo/camaras/C-1", json={"locality": "X"}),
        client.delete("/catalogo/camaras/C-1"),
        client.post("/catalogo/sedes/", json=sede),
        client.put("/catalogo/sedes/1", json=sede),
        client.delete("/catalogo/sedes/1"),
    ):
        assert response.status_code == 403, response.text


# --------------------------------------------------------------------------
# Planilla para planificar
# --------------------------------------------------------------------------


def test_la_planilla_del_catalogo_entra_al_planificador_sin_mapear(client, municipios) -> None:
    """Ida y vuelta: catálogo → planilla → /upload-excel/ → /process/, con los IDs intactos."""
    _importar(client, [_fila("00123", -24.75, -65.45), *BASE])
    # Un CSV deja pasar texto que empieza con "=" tal cual (un .xlsx lo tomaría
    # como fórmula ya al armar la planilla de prueba).
    csv = pd.DataFrame([_fila("C-F", -24.75, -65.45, "=1+1")]).to_csv(index=False).encode()
    response = client.post(
        "/catalogo/importar/", files={"file": ("f.csv", csv, "text/csv")}, data=MAPEO
    )
    assert response.status_code == 200, response.text

    response = client.post("/catalogo/planilla", json={"ids": [" C-3 ", "00123", "C-3", "C-F"]})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MIME
    assert 'filename="catalogo-3-camaras-' in response.headers["content-disposition"]
    contenido = response.content

    hoja = load_workbook(io.BytesIO(contenido)).worksheets[0]
    assert hoja.cell(row=3, column=1).value == "00123"
    assert hoja.cell(row=4, column=4).value == "=1+1"
    assert hoja.cell(row=4, column=4).data_type == "s"  # texto, no fórmula
    assert hoja.cell(row=2, column=7).value == "Cerrillos"

    subida = client.post("/upload-excel/", files={"file": ("plan.xlsx", contenido, XLSX_MIME)})
    sugerido = subida.json()["suggested_mapping"]
    assert sugerido["id"] == "ID de la cámara"
    assert (sugerido["lat"], sugerido["lon"]) == ("Latitud", "Longitud")
    assert sugerido["label"] == "Descripción"
    assert sugerido["node"] == "Nodo preliminar"
    assert sugerido["observation"] == "Observación"
    assert sugerido["done"] is None

    procesado = client.post(
        "/process/",
        files={"file": ("plan.xlsx", contenido, XLSX_MIME)},
        data={
            "col_id": sugerido["id"], "col_lat": sugerido["lat"], "col_lon": sugerido["lon"],
            "col_label": sugerido["label"], "eps_km": "5",
        },
    )
    assert procesado.status_code == 200, procesado.text
    assert [camara["id"] for camara in procesado.json()["cameras"]] == ["C-3", "00123", "C-F"]


def test_la_planilla_rechaza_ids_que_no_estan(client) -> None:
    _importar(client, BASE)
    response = client.post("/catalogo/planilla", json={"ids": ["C-1", "X-1", "X-2"]})
    assert response.status_code == 422
    assert "2 ID(s) no están en el catálogo: X-1, X-2" in response.json()["detail"]

    assert client.post("/catalogo/planilla", json={"ids": []}).status_code == 422
    assert client.post("/catalogo/planilla", json={"ids": ["  "]}).status_code == 422


# --------------------------------------------------------------------------
# Cruce con el registro
# --------------------------------------------------------------------------


def _guardar_plan(client: TestClient, jornadas: list[tuple[str, list[str]]]) -> dict:
    routes = [
        {
            "date": fecha,
            "cluster_id": 0,
            "day": dia,
            "start_lat": -24.78,
            "start_lon": -65.41,
            "distance_m": 10_000,
            "duration_s": 3_600,
            "stops": [{"camera_id": camara, "lat": -24.75, "lon": -65.45} for camara in camaras],
        }
        for dia, (fecha, camaras) in enumerate(jornadas, start=1)
    ]
    response = client.post("/registro/planes/", json={"routes": routes})
    assert response.status_code == 201, response.text
    return response.json()


def _marcar(client: TestClient, camara: str, estado: str) -> None:
    tareas = client.get("/registro/tareas/", params={"q": camara}).json()
    for tarea in tareas:
        if tarea["camera_id"] == camara:
            response = client.patch(f"/registro/tareas/{tarea['id']}", json={"status": estado})
            assert response.status_code == 200


def test_cada_camara_trae_su_ultima_visita_y_estado(client) -> None:
    _importar(client, BASE)
    _guardar_plan(client, [("2026-09-01", ["C-1", "C-2"])])
    _marcar(client, "C-1", "realizada")
    _marcar(client, "C-2", "no_realizada")
    _guardar_plan(client, [("2026-09-20", ["C-2"])])  # C-2 se reprograma

    camaras = _camaras(client)
    assert camaras["C-1"]["visits"] == 1
    assert camaras["C-1"]["last_visit"] == "2026-09-01"
    assert camaras["C-1"]["last_status"] == "realizada"
    assert camaras["C-2"]["visits"] == 0
    assert camaras["C-2"]["last_status"] == "pendiente"  # la del plan más nuevo
    assert camaras["C-2"]["last_planned"] == "2026-09-20"
    assert camaras["C-3"]["last_status"] is None


# --------------------------------------------------------------------------
# Sedes
# --------------------------------------------------------------------------


def test_sedes_alta_cambio_y_baja(client) -> None:
    creada = client.post("/catalogo/sedes/", json={"name": " Base Norte ", "lat": -24.7, "lon": -65.4})
    assert creada.status_code == 201
    sede = creada.json()
    assert sede["name"] == "Base Norte" and isinstance(sede["id"], int)

    client.post("/catalogo/sedes/", json={"name": "Anexo", "lat": -24.8, "lon": -65.5})
    assert [s["name"] for s in client.get("/catalogo/sedes/").json()] == ["Anexo", "Base Norte"]

    repetida = client.post("/catalogo/sedes/", json={"name": "base norte", "lat": 0, "lon": 0})
    assert repetida.status_code == 409

    cambiada = client.put(f"/catalogo/sedes/{sede['id']}", json={"name": "Base Centro", "lat": -24.78, "lon": -65.41})
    assert cambiada.status_code == 200 and cambiada.json()["name"] == "Base Centro"
    assert client.put(f"/catalogo/sedes/{sede['id']}", json={"name": "Anexo", "lat": 0, "lon": 0}).status_code == 409
    assert client.put("/catalogo/sedes/999", json={"name": "X", "lat": 0, "lon": 0}).status_code == 404

    assert client.delete(f"/catalogo/sedes/{sede['id']}").status_code == 204
    assert client.delete(f"/catalogo/sedes/{sede['id']}").status_code == 404


@pytest.mark.parametrize(
    "cuerpo",
    [
        {"name": "   ", "lat": 0, "lon": 0},
        {"name": "X", "lat": 95, "lon": 0},
        {"name": "X", "lat": 0, "lon": -200},
        {"name": "", "lat": 0, "lon": 0},
    ],
)
def test_sedes_invalidas(client, cuerpo) -> None:
    assert client.post("/catalogo/sedes/", json=cuerpo).status_code == 422


# --------------------------------------------------------------------------
# Clusters
# --------------------------------------------------------------------------


def test_clusters_del_catalogo_con_colores_distintos_entre_vecinos(client) -> None:
    filas = [_fila(f"A-{i}", -24.75 + i * 0.002, -65.45) for i in range(4)]
    filas += [_fila(f"B-{i}", -24.95 + i * 0.002, -65.40) for i in range(3)]
    filas += [_fila("Suelta", -26.0, -67.0)]
    _importar(client, filas)

    response = client.get("/catalogo/clusters", params={"eps_km": 2, "noise_reassign_factor": 0})
    assert response.status_code == 200
    datos = response.json()
    assert datos["eps_km"] == 2
    assert len(datos["clusters"]) == 2
    assert datos["noise_count"] == 1
    assert {c["id"]: c["cluster"] for c in datos["cameras"]}["Suelta"] == -1
    colores = {cluster["color"] for cluster in datos["clusters"]}
    assert len(colores) == 2


def test_una_camara_suelta_cercana_se_suma_al_cluster_por_defecto(client) -> None:
    """A ~1,4 radios de un grupo, la cámara suelta entra; con factor 1 queda afuera."""
    filas = [_fila(f"A-{i}", -24.75 + i * 0.002, -65.45) for i in range(4)]
    filas += [_fila("Cerca", -24.744 + 0.126, -65.45)]  # ~14 km del miembro más cercano
    _importar(client, filas)

    por_defecto = client.get("/catalogo/clusters", params={"eps_km": 10}).json()
    estricto = client.get(
        "/catalogo/clusters", params={"eps_km": 10, "noise_reassign_factor": 1}
    ).json()

    asignada = {c["id"]: c["cluster"] for c in por_defecto["cameras"]}
    suelta = {c["id"]: c["cluster"] for c in estricto["cameras"]}
    assert asignada["Cerca"] == asignada["A-0"] != -1
    assert suelta["Cerca"] == -1


def test_clusters_se_reutilizan_mientras_no_cambie_el_catalogo(client, monkeypatch) -> None:
    """El Inicio los pide cada vez que se abre: DBSCAN sobre miles de cámaras
    tarda, y el resultado sólo depende de las cámaras, las sedes y los
    parámetros."""
    catalogo.limpiar_cache_clusters()
    corridas = []
    original = catalogo.run_dbscan
    monkeypatch.setattr(catalogo, "run_dbscan", lambda *a, **k: corridas.append(1) or original(*a, **k))
    _importar(client, BASE)
    parametros = {"eps_km": 5, "por_sede": False}

    primero = client.get("/catalogo/clusters", params=parametros).json()
    segundo = client.get("/catalogo/clusters", params=parametros).json()
    otro_radio = client.get("/catalogo/clusters", params={**parametros, "eps_km": 6}).json()
    _importar(client, [_fila("C-1", -24.70, -65.45)], nombre="mueve.xlsx")  # se movió
    movido = client.get("/catalogo/clusters", params=parametros).json()

    assert primero == segundo
    assert otro_radio["eps_km"] == 6
    assert len(corridas) == 3  # primero, otro radio y después de mover una cámara
    assert movido["cameras"]


def test_clusters_con_el_catalogo_vacio(client) -> None:
    datos = client.get("/catalogo/clusters").json()
    assert datos == {
        "eps_km": 20.0, "by_depot": True, "depot_max_km": 60.0,
        "cameras": [], "clusters": [], "noise_count": 0,
    }


def _cluster(cluster_id: int, lon: float, size: int = 3) -> dict:
    return {"id": cluster_id, "size": size, "centroid_lat": -24.8, "centroid_lon": lon, "radius_km": 0.5}


def test_con_mas_clusters_que_colores_se_repiten_los_mas_lejanos() -> None:
    # Tres en fila, cada uno a ~10 km del siguiente, con dos colores.
    clusters = [_cluster(0, -65.5), _cluster(1, -65.4), _cluster(2, -65.3)]
    colores = catalogo.assign_colors(clusters, 2)
    assert colores[0] != colores[1] and colores[1] != colores[2]
    assert colores[0] == colores[2]
    assert catalogo.assign_colors([], 8) == []


def test_con_colores_de_sobra_no_se_repite_ninguno() -> None:
    clusters = [_cluster(i, -65.5 + i * 0.1, size=10 - i) for i in range(6)]
    assert len(set(catalogo.assign_colors(clusters, 8))) == 6


# --------------------------------------------------------------------------
# Resumen del Inicio
# --------------------------------------------------------------------------


def test_resumen_del_inicio(client, municipios) -> None:
    hoy = dt.date.today()
    _importar(client, BASE + [_fila("C-5", -24.76, -65.46)])
    client.post("/catalogo/sedes/", json={"name": "Centro", "lat": -24.78, "lon": -65.41})
    client.post("/catalogo/sedes/", json={"name": "Sur", "lat": -24.96, "lon": -65.40})

    hace_10 = (hoy - dt.timedelta(days=10)).isoformat()
    hace_100 = (hoy - dt.timedelta(days=100)).isoformat()
    _guardar_plan(client, [(hace_100, ["C-3"]), (hace_10, ["C-1", "C-2", "X-fuera"])])
    _marcar(client, "C-1", "realizada")
    _marcar(client, "C-3", "realizada")

    resumen = client.get("/catalogo/resumen").json()

    assert resumen["today"] == hoy.isoformat()
    assert resumen["camera_count"] == 4
    assert resumen["locality_count"] == 3
    assert resumen["municipalities_loaded"] is True
    salta, *_ = resumen["localities"]
    assert salta == {
        "name": "Salta", "cameras": 2, "visited": 1, "pending": 0,
        "last_visit": hace_10,
    }
    enclave = next(item for item in resumen["localities"] if item["name"] == "Enclave")
    assert enclave["pending"] == 1  # C-2 sigue pendiente

    assert resumen["visit_age"] == {"never": 2, "within_30": 1, "within_90": 0, "older": 1}
    assert resumen["plan_count"] == 1 and resumen["route_count"] == 2
    assert resumen["distance_m"] == 20_000
    assert resumen["tasks"] == {
        "total": 4, "done": 2, "pending": 2, "not_done": 0, "rescheduled": 0,
    }
    meses = resumen["months"]
    assert len(meses) == catalogo.MONTHS_BACK
    assert meses[-1]["month"] == hoy.strftime("%Y-%m")
    assert sum(mes["done"] + mes["pending"] for mes in meses) == 4

    sedes = {sede["name"]: sede for sede in resumen["depots"]}
    assert sedes["Centro"]["cameras"] == 3
    assert sedes["Sur"]["cameras"] == 1 and sedes["Sur"]["max_km"] == pytest.approx(1.1, abs=0.2)
    assert resumen["planned_outside_catalog"] == 1
    assert resumen["last_import_at"] is not None


def test_resumen_con_jornadas_futuras_y_sin_datos(client) -> None:
    vacio = client.get("/catalogo/resumen").json()
    assert vacio["camera_count"] == 0 and vacio["localities"] == [] and vacio["depots"] == []
    assert vacio["last_import_at"] is None

    hoy = dt.date.today()
    en_dos_meses = (hoy.replace(day=1) + dt.timedelta(days=62)).isoformat()
    _guardar_plan(client, [(en_dos_meses, ["C-1"])])
    meses = client.get("/catalogo/resumen").json()["months"]
    assert meses[-1]["month"] == en_dos_meses[:7]
    assert meses[-1]["pending"] == 1
    assert len(meses) == catalogo.MONTHS_BACK + 2


def test_sede_sin_camaras_en_el_resumen(client) -> None:
    client.post("/catalogo/sedes/", json={"name": "Sola", "lat": -24.78, "lon": -65.41})
    sede = client.get("/catalogo/resumen").json()["depots"][0]
    assert sede == {"id": sede["id"], "name": "Sola", "cameras": 0, "average_km": None, "max_km": None}
