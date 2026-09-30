"""Tests del Excel de seguimiento que se entrega a los técnicos."""

import datetime as dt
import io

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.schemas import ExportDay, ExportRequest, ExportStop
from app.services.export import FOLLOW_UP_COLUMNS, build_plan_workbook
from app.services.ingest import extract_points, split_done, suggest_mapping

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Columnas que pidió el usuario, en este orden, antes de las de replanificación.
COLUMNAS_PEDIDAS = [
    "ID de la cámara",
    "Fecha de planificación",
    "Orden",
    "Realizado",
    "Observación",
    "Nodo al cual se migró",
    "Nodo preliminar",
]


def _parada(camera_id: str, lat: float = -24.83, **extra: object) -> ExportStop:
    return ExportStop(camera_id=camera_id, lat=lat, lon=-65.41, **extra)


def _plan() -> ExportRequest:
    """Dos días cargados fuera de orden, para verificar que se ordenan por fecha."""
    return ExportRequest(
        days=[
            ExportDay(
                date=dt.date(2026, 10, 2), cluster_id=0, day=2, distance_m=12500,
                duration_s=5400, stops=[_parada("C-3"), _parada("C-4")],
            ),
            ExportDay(
                date=dt.date(2026, 10, 1), cluster_id=0, day=1, distance_m=8000,
                duration_s=3600, start_name="Base",
                stops=[
                    _parada("C-1", node="ATOCHA", observation="Sin acceso"),
                    _parada("C-2"),
                ],
            ),
        ]
    )


def _hoja(contenido: bytes, nombre: str | None = None):
    libro = load_workbook(io.BytesIO(contenido))
    return libro[nombre] if nombre else libro.worksheets[0]


def test_la_primera_hoja_tiene_las_columnas_pedidas_en_orden() -> None:
    """Primera hoja = la que lee la app al reimportar; encabezados en la fila 1."""
    hoja = _hoja(build_plan_workbook(_plan()))

    encabezados = [celda.value for celda in hoja[1]]
    assert encabezados[: len(COLUMNAS_PEDIDAS)] == COLUMNAS_PEDIDAS
    assert encabezados == FOLLOW_UP_COLUMNS
    assert "Latitud" in encabezados and "Longitud" in encabezados


def test_filas_ordenadas_por_fecha_y_orden_de_visita() -> None:
    """Una fila por cámara, por fecha y luego por orden del recorrido."""
    hoja = _hoja(build_plan_workbook(_plan()))

    filas = [(fila[0], fila[1], fila[2]) for fila in hoja.iter_rows(min_row=2, values_only=True)]
    assert filas == [
        ("C-1", dt.datetime(2026, 10, 1), 1),
        ("C-2", dt.datetime(2026, 10, 1), 2),
        ("C-3", dt.datetime(2026, 10, 2), 1),
        ("C-4", dt.datetime(2026, 10, 2), 2),
    ]


def test_la_fecha_es_una_fecha_de_excel_no_texto() -> None:
    """Fecha real con formato dd/mm/aaaa: se puede filtrar y ordenar en Excel."""
    hoja = _hoja(build_plan_workbook(_plan()))

    assert hoja["B2"].is_date
    assert hoja["B2"].number_format == "DD/MM/YYYY"


def test_realizado_se_completa_con_una_lista_si_no() -> None:
    """La columna Realizado ofrece Sí/No en un desplegable."""
    hoja = _hoja(build_plan_workbook(_plan()))

    validaciones = hoja.data_validations.dataValidation
    assert len(validaciones) == 1
    assert validaciones[0].formula1 == '"Sí,No"'
    assert "D2" in validaciones[0].sqref


def test_arrastra_nodo_y_observacion() -> None:
    """El nodo preliminar y la observación previa llegan precargados."""
    hoja = _hoja(build_plan_workbook(_plan()))

    fila = next(hoja.iter_rows(min_row=2, max_row=2, values_only=True))
    registro = dict(zip(FOLLOW_UP_COLUMNS, fila, strict=True))
    assert registro["Nodo preliminar"] == "ATOCHA"
    assert registro["Observación"] == "Sin acceso"
    assert registro["Nodo al cual se migró"] is None
    assert registro["Realizado"] is None


def test_un_texto_que_empieza_con_igual_no_se_vuelve_formula() -> None:
    """Un ID u observación con '=' al inicio queda como texto, no se ejecuta."""
    plan = ExportRequest(
        days=[
            ExportDay(
                date=dt.date(2026, 10, 1), cluster_id=0, day=1, distance_m=0, duration_s=0,
                stops=[_parada("=HYPERLINK(\"http://x\")", observation="=1+1")],
            )
        ]
    )

    hoja = _hoja(build_plan_workbook(plan))

    assert hoja["A2"].data_type == "s"
    assert hoja["A2"].value == '=HYPERLINK("http://x")'
    assert hoja["E2"].data_type == "s"


def test_resumen_por_dia() -> None:
    """La hoja Resumen trae un renglón por día con cámaras, km y duración."""
    hoja = _hoja(build_plan_workbook(_plan()), "Resumen")

    filas = list(hoja.iter_rows(min_row=2, max_row=3, values_only=True))
    assert [fila[1] for fila in filas] == [1, 2]  # día
    assert [fila[4] for fila in filas] == [2, 2]  # cámaras
    assert filas[0][5] == pytest.approx(8.0)  # km
    assert filas[0][3] == "Base"


def test_ida_y_vuelta_el_excel_completado_se_puede_reimportar() -> None:
    """Exportar, tildar una como realizada y volver a leerlo con la ingesta."""
    contenido = build_plan_workbook(_plan())
    frame = pd.read_excel(io.BytesIO(contenido))
    # La columna llega vacía (float NaN); pandas 3 no deja meter texto sin convertir.
    frame["Realizado"] = frame["Realizado"].astype(object)
    frame.loc[frame["ID de la cámara"] == "C-1", "Realizado"] = "Sí"

    mapeo = suggest_mapping(list(frame.columns))
    pendientes, hechas = split_done(frame, mapeo["done"])
    puntos, descartadas = extract_points(
        pendientes, mapeo["id"], mapeo["lat"], mapeo["lon"],
        col_node=mapeo["node"], col_obs=mapeo["observation"],
    )

    assert hechas == 1
    assert descartadas == []
    assert puntos["id"].tolist() == ["C-2", "C-3", "C-4"]


# --------------------------------------------------------------------------
# POST /export/
# --------------------------------------------------------------------------


def test_endpoint_devuelve_el_xlsx_como_descarga() -> None:
    """El endpoint responde el archivo con nombre sugerido por la primera fecha."""
    response = TestClient(app).post("/export/", json=_plan().model_dump(mode="json"))

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MIME
    assert "plan-recorridos-2026-10-01.xlsx" in response.headers["content-disposition"]
    assert [c.value for c in _hoja(response.content)[1]] == FOLLOW_UP_COLUMNS


def test_endpoint_rechaza_un_plan_vacio() -> None:
    """Sin días no hay nada que exportar: 422 de validación."""
    response = TestClient(app).post("/export/", json={"days": []})

    assert response.status_code == 422
