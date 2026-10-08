"""Tests del registro de recorridos: base SQLite, endpoints y carga de seguimientos."""

import io
import sqlite3

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import config, database
from app.main import app
from app.services.export import REGISTRY_SHEET, read_plan_id
from app.services.ingest import (
    extract_points,
    follow_up_mapping,
    reported_status,
    split_done,
    suggest_mapping,
)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def client() -> TestClient:
    return TestClient(app)


def _parada(camera_id: str, lat: float = -24.83, **extra: object) -> dict:
    return {"camera_id": camera_id, "lat": lat, "lon": -65.41, **extra}


def _jornada(fecha: str, dia: int, camaras: list[str], cluster: int = 0, **extra: object) -> dict:
    return {
        "date": fecha,
        "cluster_id": cluster,
        "day": dia,
        "start_name": "Base Centro",
        "start_lat": -24.78,
        "start_lon": -65.41,
        "distance_m": 12_000,
        "duration_s": 5_400,
        "geometry": [[-24.78, -65.41], [-24.83, -65.41]],
        "stops": [_parada(camara) for camara in camaras],
        **extra,
    }


def _plan(*jornadas: dict, **extra: object) -> dict:
    return {
        "source_file": "camaras.xlsx",
        "provider": "osrm",
        "is_road_network": True,
        "routes": list(jornadas)
        or [
            _jornada("2026-10-01", 1, ["C-1", "C-2"]),
            _jornada("2026-10-02", 2, ["C-3"]),
        ],
        **extra,
    }


def _guardar(client: TestClient, plan: dict | None = None) -> dict:
    response = client.post("/registro/planes/", json=plan or _plan())
    assert response.status_code == 201, response.text
    return response.json()


def _tareas(client: TestClient, **params: object) -> dict[str, dict]:
    response = client.get("/registro/tareas/", params=params)
    assert response.status_code == 200, response.text
    return {(t["camera_id"], t["plan_id"]): t for t in response.json()}


def _estados(client: TestClient, **params: object) -> dict[tuple[str, int], str]:
    return {clave: tarea["status"] for clave, tarea in _tareas(client, **params).items()}


def _exportar(client: TestClient, plan: dict, plan_id: int | None) -> bytes:
    """Excel de seguimiento del plan, como lo pide el frontend."""
    dias = [
        {
            "date": jornada["date"],
            "cluster_id": jornada["cluster_id"],
            "day": jornada["day"],
            "start_name": jornada["start_name"],
            "distance_m": jornada["distance_m"],
            "duration_s": jornada["duration_s"],
            "stops": jornada["stops"],
        }
        for jornada in plan["routes"]
    ]
    response = client.post("/export/", json={"days": dias, "plan_id": plan_id})
    assert response.status_code == 200, response.text
    return response.content


def _completar(contenido: bytes, cambios: dict[str, dict[str, str]]) -> bytes:
    """Simula a los técnicos: completa Realizado, Observación y nodo por cámara."""
    libro = load_workbook(io.BytesIO(contenido))
    hoja = libro.worksheets[0]
    columnas = {celda.value: celda.column for celda in hoja[1]}
    for fila in hoja.iter_rows(min_row=2):
        camara = fila[0].value
        for columna, valor in cambios.get(camara, {}).items():
            hoja.cell(row=fila[0].row, column=columnas[columna], value=valor)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _cargar(client: TestClient, contenido: bytes, nombre: str = "seguimiento.xlsx") -> dict:
    mime = XLSX_MIME if nombre.endswith(".xlsx") else "text/csv"
    response = client.post(
        "/registro/seguimiento/", files={"file": (nombre, contenido, mime)}
    )
    assert response.status_code == 200, response.text
    return response.json()


# --------------------------------------------------------------------------
# La base
# --------------------------------------------------------------------------


def test_la_base_se_crea_con_tablas_vistas_y_version() -> None:
    """Arrancar alcanza para tener la base lista: tablas, vistas y versión."""
    conn = database.connect()
    try:
        objetos = {
            fila["name"]
            for fila in conn.execute("SELECT name FROM sqlite_master WHERE type IN ('table', 'view')")
        }
        version = conn.execute("PRAGMA user_version").fetchone()[0]
    finally:
        conn.close()

    assert {"planes", "recorridos", "paradas", "cargas_seguimiento"} <= objetos
    assert {"camaras", "sedes", "cargas_catalogo"} <= objetos
    assert {"v_paradas", "v_recorridos", "v_planes"} <= objetos
    assert version == database.SCHEMA_VERSION


def test_si_se_borra_el_archivo_se_vuelve_a_crear(tmp_path) -> None:
    """Borrar la base con el backend andando no deja requests sin tablas."""
    ruta = tmp_path / "otra.db"
    database.connect(str(ruta)).close()
    ruta.unlink()

    conn = database.connect(str(ruta))
    try:
        assert conn.execute("SELECT COUNT(*) FROM planes").fetchone()[0] == 0
    finally:
        conn.close()


def test_borrar_un_plan_borra_sus_jornadas_y_tareas(client: TestClient) -> None:
    """Las claves foráneas están activas: el borrado en cascada funciona."""
    plan = _guardar(client)

    assert client.delete(f"/registro/planes/{plan['id']}").status_code == 204

    conn = database.connect()
    try:
        assert conn.execute("SELECT COUNT(*) FROM recorridos").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM paradas").fetchone()[0] == 0
    finally:
        conn.close()


# --------------------------------------------------------------------------
# Planes
# --------------------------------------------------------------------------


def test_guardar_un_plan_devuelve_su_avance(client: TestClient) -> None:
    plan = _guardar(client)

    assert plan["name"] == "Plan del 01/10 al 02/10/2026"
    assert plan["route_count"] == 2
    assert plan["total"] == plan["pending"] == 3
    assert plan["done"] == plan["not_done"] == plan["rescheduled"] == 0
    assert plan["date_from"] == "2026-10-01"
    assert plan["date_to"] == "2026-10-02"
    assert plan["distance_m"] == pytest.approx(24_000)
    assert plan["is_road_network"] is True


def test_un_plan_de_un_solo_dia_se_nombra_con_esa_fecha(client: TestClient) -> None:
    plan = _guardar(client, _plan(_jornada("2026-10-05", 1, ["C-1"])))

    assert plan["name"] == "Plan del 05/10/2026"


def test_un_plan_que_repite_una_jornada_se_rechaza(client: TestClient) -> None:
    repetido = _plan(_jornada("2026-10-01", 1, ["C-1"]), _jornada("2026-10-02", 1, ["C-2"]))

    response = client.post("/registro/planes/", json=repetido)

    assert response.status_code == 422
    assert "jornada 1 del cluster 0" in response.json()["detail"]
    assert client.get("/registro/planes/").json() == []  # no quedó a medias


def test_renombrar_y_anotar_un_plan(client: TestClient) -> None:
    plan = _guardar(client)

    response = client.patch(
        f"/registro/planes/{plan['id']}", json={"name": "Zona norte", "notes": "Llevar escalera"}
    )

    assert response.status_code == 200
    assert response.json()["name"] == "Zona norte"
    assert response.json()["notes"] == "Llevar escalera"


def test_plan_inexistente_da_404(client: TestClient) -> None:
    assert client.patch("/registro/planes/99", json={"name": "x"}).status_code == 404
    assert client.delete("/registro/planes/99").status_code == 404
    assert client.put("/registro/planes/99", json=_plan()).status_code == 404


def test_reemplazar_un_plan_sin_seguimiento_no_lo_duplica(client: TestClient) -> None:
    """Volver a exportar con otra fecha actualiza el mismo plan."""
    plan = _guardar(client)
    movido = _plan(_jornada("2026-10-08", 1, ["C-1", "C-2"]), _jornada("2026-10-09", 2, ["C-3"]))

    response = client.put(f"/registro/planes/{plan['id']}", json=movido)

    assert response.status_code == 200
    assert response.json()["date_from"] == "2026-10-08"
    # El nombre automático sigue a las fechas nuevas.
    assert response.json()["name"] == "Plan del 08/10 al 09/10/2026"
    assert len(client.get("/registro/planes/").json()) == 1
    assert {r["date"] for r in client.get("/registro/recorridos/").json()} == {
        "2026-10-08", "2026-10-09",
    }


def test_reemplazar_respeta_un_nombre_puesto_a_mano(client: TestClient) -> None:
    plan = _guardar(client)
    client.patch(f"/registro/planes/{plan['id']}", json={"name": "Zona norte"})

    response = client.put(
        f"/registro/planes/{plan['id']}", json=_plan(_jornada("2026-10-20", 1, ["C-1"]))
    )

    assert response.json()["name"] == "Zona norte"


def test_reemplazar_con_una_jornada_repetida_no_toca_el_plan(client: TestClient) -> None:
    plan = _guardar(client)
    repetido = _plan(_jornada("2026-10-01", 1, ["C-1"]), _jornada("2026-10-02", 1, ["C-2"]))

    response = client.put(f"/registro/planes/{plan['id']}", json=repetido)

    assert response.status_code == 422
    assert client.get("/registro/planes/").json()[0]["total"] == 3  # rollback completo


def test_no_se_reemplaza_un_plan_con_seguimiento(client: TestClient) -> None:
    """Reemplazarlo borraría lo que informaron los técnicos."""
    plan = _guardar(client)
    tarea = next(iter(_tareas(client).values()))
    client.patch(f"/registro/tareas/{tarea['id']}", json={"status": "realizada"})

    response = client.put(f"/registro/planes/{plan['id']}", json=_plan())

    assert response.status_code == 409
    assert "plan nuevo" in response.json()["detail"]


# --------------------------------------------------------------------------
# Recorridos y filtros
# --------------------------------------------------------------------------


def test_recorridos_ordenados_del_mas_reciente(client: TestClient) -> None:
    _guardar(client)

    fechas = [r["date"] for r in client.get("/registro/recorridos/").json()]

    assert fechas == ["2026-10-02", "2026-10-01"]


def test_filtros_por_fecha_plan_y_texto(client: TestClient) -> None:
    primero = _guardar(client)
    segundo = _guardar(
        client,
        _plan(_jornada("2026-10-10", 1, ["Z-9"], start_name="Base Norte")),
    )

    def fechas(**params: object) -> list[str]:
        response = client.get("/registro/recorridos/", params=params)
        assert response.status_code == 200, response.text
        return [r["date"] for r in response.json()]

    assert fechas(desde="2026-10-02") == ["2026-10-10", "2026-10-02"]
    assert fechas(hasta="2026-10-01") == ["2026-10-01"]
    assert fechas(plan_id=primero["id"]) == ["2026-10-02", "2026-10-01"]
    assert fechas(plan_id=segundo["id"]) == ["2026-10-10"]
    # Una cámara encuentra la jornada en que se visita; la salida también busca.
    assert fechas(q="c-3") == ["2026-10-02"]
    assert fechas(q="norte") == ["2026-10-10"]
    tarea = _tareas(client)[("C-1", primero["id"])]
    client.patch(f"/registro/tareas/{tarea['id']}", json={"observation": "Poste caído"})
    assert fechas(q="poste") == ["2026-10-01"]


def test_la_busqueda_toma_los_comodines_como_texto(client: TestClient) -> None:
    """Un "%" escrito por el usuario no es un comodín que traiga todo."""
    _guardar(client)

    assert client.get("/registro/recorridos/", params={"q": "%"}).json() == []


def test_rango_de_fechas_invertido_da_422(client: TestClient) -> None:
    response = client.get("/registro/recorridos/", params={"desde": "2026-10-05", "hasta": "2026-10-01"})

    assert response.status_code == 422


def test_detalle_trae_paradas_y_polilinea_en_el_orden_pedido(client: TestClient) -> None:
    _guardar(client)
    ids = [r["id"] for r in client.get("/registro/recorridos/").json()]

    response = client.get("/registro/recorridos/detalle", params={"ids": [ids[1], ids[0], 999]})

    assert response.status_code == 200
    detalle = response.json()
    assert [r["id"] for r in detalle] == [ids[1], ids[0]]  # el inexistente se omite
    assert detalle[0]["geometry"] == [[-24.78, -65.41], [-24.83, -65.41]]
    assert [p["camera_id"] for p in detalle[0]["stops"]] == ["C-1", "C-2"]
    assert [p["order"] for p in detalle[0]["stops"]] == [1, 2]


# --------------------------------------------------------------------------
# Carga de seguimientos
# --------------------------------------------------------------------------


def test_el_excel_exportado_lleva_el_plan_en_una_hoja_oculta(client: TestClient) -> None:
    plan = _plan()
    guardado = _guardar(client, plan)

    contenido = _exportar(client, plan, guardado["id"])

    libro = load_workbook(io.BytesIO(contenido))
    assert libro.worksheets[0].title == "Recorridos"  # la que lee la ingesta
    assert libro[REGISTRY_SHEET].sheet_state == "hidden"
    assert read_plan_id("plan.xlsx", contenido) == guardado["id"]


def test_sin_plan_no_hay_hoja_oculta(client: TestClient) -> None:
    contenido = _exportar(client, _plan(), None)

    assert REGISTRY_SHEET not in load_workbook(io.BytesIO(contenido)).sheetnames
    assert read_plan_id("plan.xlsx", contenido) is None
    assert read_plan_id("plan.csv", b"a,b") is None
    assert read_plan_id("roto.xlsx", b"no es un excel") is None


def test_cargar_el_seguimiento_actualiza_las_tareas(client: TestClient) -> None:
    """El ciclo completo: guardar, exportar, completar y volver a cargar."""
    plan = _plan()
    guardado = _guardar(client, plan)
    completado = _completar(
        _exportar(client, plan, guardado["id"]),
        {
            "C-1": {"Realizado": "Sí", "Nodo al cual se migró": "NODO-7"},
            "C-2": {"Realizado": "No", "Observación": "Sin acceso al poste"},
        },
    )

    resultado = _cargar(client, completado)

    assert resultado["plan_id"] == guardado["id"]
    assert resultado["plan_name"] == guardado["name"]
    assert resultado["rows"] == resultado["matched"] == 3
    assert (resultado["done"], resultado["not_done"], resultado["no_news"]) == (1, 1, 1)
    assert resultado["updated"] == 2
    assert resultado["unmatched"] == 0
    assert resultado["previously_loaded_at"] is None

    tareas = _tareas(client)
    c1, c2, c3 = (tareas[(c, guardado["id"])] for c in ("C-1", "C-2", "C-3"))
    assert (c1["status"], c1["migrated_node"]) == ("realizada", "NODO-7")
    assert (c2["status"], c2["observation"]) == ("no_realizada", "Sin acceso al poste")
    assert c3["status"] == "pendiente"
    assert c1["verified_at"] is not None and c3["verified_at"] is None

    resumen = client.get("/registro/planes/").json()[0]
    assert (resumen["done"], resumen["not_done"], resumen["pending"]) == (1, 1, 1)


def test_cargar_dos_veces_el_mismo_archivo_no_cambia_nada(client: TestClient) -> None:
    plan = _plan()
    guardado = _guardar(client, plan)
    completado = _completar(_exportar(client, plan, guardado["id"]), {"C-1": {"Realizado": "Sí"}})
    _cargar(client, completado)

    repetido = _cargar(client, completado)

    assert repetido["updated"] == 0
    assert repetido["previously_loaded_at"] is not None
    assert len(client.get("/registro/cargas/").json()) == 2  # el historial lo registra igual


def test_una_celda_vacia_no_deshace_lo_ya_informado(client: TestClient) -> None:
    """Un seguimiento parcial posterior no vuelve a pendiente lo realizado."""
    plan = _plan()
    guardado = _guardar(client, plan)
    exportado = _exportar(client, plan, guardado["id"])
    _cargar(client, _completar(exportado, {"C-1": {"Realizado": "Sí"}}))

    _cargar(client, _completar(exportado, {"C-2": {"Realizado": "Sí"}}), "segundo.xlsx")

    estados = _estados(client)
    assert estados[("C-1", guardado["id"])] == "realizada"
    assert estados[("C-2", guardado["id"])] == "realizada"


def test_sin_hoja_oculta_se_cruza_por_camara_y_fecha(client: TestClient) -> None:
    """Un CSV (o un Excel rehecho) se cruza igual, y gana el plan más nuevo."""
    viejo = _guardar(client, _plan(_jornada("2026-10-01", 1, ["C-1"])))
    nuevo = _guardar(client, _plan(_jornada("2026-10-01", 1, ["C-1"])))
    csv = "ID de la cámara,Fecha de planificación,Realizado\nC-1,01/10/2026,Sí\n".encode()

    resultado = _cargar(client, csv, "seguimiento.csv")

    assert resultado["plan_id"] == nuevo["id"]
    estados = _estados(client)
    assert estados[("C-1", nuevo["id"])] == "realizada"
    assert estados[("C-1", viejo["id"])] == "reprogramada"


def test_si_la_fecha_no_coincide_toma_la_tarea_mas_reciente(client: TestClient) -> None:
    """El técnico pudo haber corregido la fecha: se cruza por cámara."""
    guardado = _guardar(client, _plan(_jornada("2026-10-01", 1, ["C-1"])))
    csv = "ID de la cámara,Fecha de planificación,Realizado\nC-1,2026-10-03,x\n".encode()

    resultado = _cargar(client, csv, "seguimiento.csv")

    assert resultado["matched"] == 1
    assert _estados(client)[("C-1", guardado["id"])] == "realizada"


def test_las_filas_sin_tarea_se_informan(client: TestClient) -> None:
    """Filas nuevas que agregó el usuario: no hay tarea guardada que actualizar."""
    _guardar(client, _plan(_jornada("2026-10-01", 1, ["C-1"])))
    csv = (
        "ID de la cámara,Realizado\nC-1,Sí\nNUEVA-1,\nNUEVA-2,Sí\n,Sí\n"
    ).encode()

    resultado = _cargar(client, csv, "seguimiento.csv")

    assert resultado["rows"] == 3  # la fila sin ID no cuenta
    assert resultado["unmatched"] == 2
    assert resultado["unmatched_ids"] == ["NUEVA-1", "NUEVA-2"]


def test_un_archivo_que_no_es_seguimiento_da_422(client: TestClient) -> None:
    csv = "id,latitud,longitud\nC-1,-24.8,-65.4\n".encode()

    response = client.post(
        "/registro/seguimiento/", files={"file": ("planilla.csv", csv, "text/csv")}
    )

    assert response.status_code == 422
    assert "Realizado" in response.json()["detail"]


def test_un_archivo_ilegible_da_422(client: TestClient) -> None:
    response = client.post(
        "/registro/seguimiento/", files={"file": ("roto.xlsx", b"no es un excel", XLSX_MIME)}
    )

    assert response.status_code == 422


def test_si_el_plan_de_la_hoja_oculta_se_borro_se_cruza_sin_el(client: TestClient) -> None:
    plan = _plan()
    borrado = _guardar(client, plan)
    exportado = _exportar(client, plan, borrado["id"])
    client.delete(f"/registro/planes/{borrado['id']}")
    vigente = _guardar(client, plan)

    resultado = _cargar(client, _completar(exportado, {"C-1": {"Realizado": "Sí"}}))

    assert resultado["plan_id"] == vigente["id"]
    assert _estados(client)[("C-1", vigente["id"])] == "realizada"


# --------------------------------------------------------------------------
# Tareas, reprogramación y próximos recorridos
# --------------------------------------------------------------------------


def test_lo_que_vuelve_a_planificarse_queda_reprogramado(client: TestClient) -> None:
    """La pendiente del plan viejo ya no "falta": está en el plan nuevo."""
    viejo = _guardar(client)
    hecha = _tareas(client)[("C-1", viejo["id"])]
    client.patch(f"/registro/tareas/{hecha['id']}", json={"status": "realizada"})

    nuevo = _guardar(client, _plan(_jornada("2026-10-12", 1, ["C-1", "C-2", "C-3"])))

    estados = _estados(client)
    assert estados[("C-1", viejo["id"])] == "realizada"  # lo hecho no se reprograma
    assert estados[("C-2", viejo["id"])] == "reprogramada"
    assert estados[("C-3", viejo["id"])] == "reprogramada"
    faltan = _tareas(client, estado="faltan")
    assert set(faltan) == {("C-1", nuevo["id"]), ("C-2", nuevo["id"]), ("C-3", nuevo["id"])}

    resumen_viejo = client.get("/registro/planes/").json()[1]
    assert (resumen_viejo["done"], resumen_viejo["rescheduled"]) == (1, 2)

    # Borrar el plan nuevo devuelve las tareas viejas a pendientes.
    client.delete(f"/registro/planes/{nuevo['id']}")
    assert _estados(client)[("C-2", viejo["id"])] == "pendiente"


def test_filtro_por_estado_admite_varios(client: TestClient) -> None:
    guardado = _guardar(client)
    tareas = _tareas(client)
    client.patch(f"/registro/tareas/{tareas[('C-1', guardado['id'])]['id']}", json={"status": "realizada"})
    client.patch(f"/registro/tareas/{tareas[('C-2', guardado['id'])]['id']}", json={"status": "no_realizada"})

    assert set(_tareas(client, estado="realizada")) == {("C-1", guardado["id"])}
    assert set(_tareas(client, estado="faltan")) == {("C-2", guardado["id"]), ("C-3", guardado["id"])}
    assert set(_tareas(client, estado=["realizada", "no_realizada"])) == {
        ("C-1", guardado["id"]), ("C-2", guardado["id"]),
    }
    assert client.get("/registro/tareas/", params={"estado": "otra"}).status_code == 422


def test_corregir_una_tarea_a_mano(client: TestClient) -> None:
    guardado = _guardar(client)
    tarea = _tareas(client)[("C-1", guardado["id"])]

    response = client.patch(
        f"/registro/tareas/{tarea['id']}",
        json={"status": "no_realizada", "observation": "  Cámara robada  ", "migrated_node": ""},
    )

    assert response.status_code == 200
    corregida = response.json()
    assert corregida["status"] == "no_realizada"
    assert corregida["observation"] == "Cámara robada"
    assert corregida["migrated_node"] is None
    assert corregida["verified_at"] is not None


def _correcciones() -> list[dict]:
    conn = database.connect()
    try:
        return [dict(row) for row in conn.execute("SELECT * FROM correcciones ORDER BY id")]
    finally:
        conn.close()


def test_un_operador_no_puede_corregir_tareas(client: TestClient, como_operador) -> None:
    guardado = _guardar(client)
    tarea = _tareas(client)[("C-1", guardado["id"])]

    response = client.patch(f"/registro/tareas/{tarea['id']}", json={"status": "realizada"})

    assert response.status_code == 403
    assert _tareas(client)[("C-1", guardado["id"])]["status"] == "pendiente"
    assert _correcciones() == []


def test_cada_correccion_registra_quien_la_hizo(client: TestClient) -> None:
    guardado = _guardar(client)
    tarea = _tareas(client)[("C-1", guardado["id"])]

    corregida = client.patch(
        f"/registro/tareas/{tarea['id']}",
        json={"status": "realizada", "observation": "Cambio de fuente", "migrated_node": "N-7"},
    ).json()

    assert corregida["corrected_by"] == "Pruebas"
    assert corregida["corrected_at"] is not None
    filas = _correcciones()
    assert {(f["campo"], f["valor_anterior"], f["valor_nuevo"]) for f in filas} == {
        ("estado", "pendiente", "realizada"),
        ("observacion", None, "Cambio de fuente"),
        ("nodo_migrado", None, "N-7"),
    }
    assert {(f["usuario"], f["nombre"], f["parada_id"], f["camara_id"]) for f in filas} == {
        ("pruebas", "Pruebas", tarea["id"], "C-1"),
    }


def test_una_correccion_sin_cambios_no_queda_registrada(client: TestClient) -> None:
    guardado = _guardar(client)
    tarea = _tareas(client)[("C-1", guardado["id"])]

    sin_cambios = client.patch(f"/registro/tareas/{tarea['id']}", json={"status": "pendiente"}).json()

    assert _correcciones() == []
    assert sin_cambios["corrected_by"] is None and sin_cambios["verified_at"] is None


def test_el_historial_de_correcciones_sobrevive_al_borrar_el_plan(client: TestClient) -> None:
    guardado = _guardar(client)
    tarea = _tareas(client)[("C-1", guardado["id"])]
    client.patch(f"/registro/tareas/{tarea['id']}", json={"status": "no_realizada"})

    assert client.delete(f"/registro/planes/{guardado['id']}").status_code == 204

    (fila,) = _correcciones()
    assert fila["parada_id"] is None
    assert (fila["camara_id"], fila["plan_id"], fila["usuario"]) == ("C-1", guardado["id"], "pruebas")


def _actividad() -> list[tuple]:
    conn = database.connect()
    try:
        return [
            (row["accion"], row["usuario"], row["plan_id"], row["plan_nombre"])
            for row in conn.execute("SELECT * FROM actividad ORDER BY id")
        ]
    finally:
        conn.close()


def test_queda_registrado_quien_crea_cambia_y_borra_cada_plan(client: TestClient) -> None:
    guardado = _guardar(client, _plan(name="Semana 1"))
    plan_id = guardado["id"]
    assert client.put(f"/registro/planes/{plan_id}", json=_plan(name="Semana 1")).status_code == 200
    assert client.patch(f"/registro/planes/{plan_id}", json={"name": "Semana uno"}).status_code == 200
    assert client.delete(f"/registro/planes/{plan_id}").status_code == 204

    assert _actividad() == [
        ("plan_creado", "pruebas", plan_id, "Semana 1"),
        ("plan_reemplazado", "pruebas", plan_id, "Semana 1"),
        ("plan_editado", "pruebas", plan_id, "Semana uno"),
        ("plan_borrado", "pruebas", plan_id, "Semana uno"),
    ]


def test_un_cambio_rechazado_no_queda_en_la_actividad(client: TestClient) -> None:
    assert client.delete("/registro/planes/999").status_code == 404
    assert _actividad() == []


def test_queda_registrado_quien_carga_un_seguimiento(client: TestClient) -> None:
    plan = _plan()
    guardado = _guardar(client, plan)
    excel = _completar(_exportar(client, plan, guardado["id"]), {"C-1": {"Realizado": "Sí"}})

    _cargar(client, excel)

    ((accion, usuario, plan_id, _),) = _actividad()[1:]
    assert (accion, usuario, plan_id) == ("seguimiento_cargado", "pruebas", guardado["id"])


def test_una_base_v3_se_migra_al_abrirla(client: TestClient) -> None:
    """Una base anterior (sin `paradas.plan_id`) se completa sola al abrirla, y
    lo reprogramado se sigue derivando igual."""
    primero = _guardar(client, _plan(_jornada("2026-10-01", 1, ["C-1", "C-2"])))
    segundo = _guardar(client, _plan(_jornada("2026-10-08", 1, ["C-1"])))
    conn = database.connect()
    for vista in database._VIEWS:
        conn.execute(f"DROP VIEW {vista}")
    conn.execute("DROP INDEX idx_paradas_camara_plan")
    conn.execute("DROP INDEX idx_paradas_plan")
    conn.execute("ALTER TABLE paradas DROP COLUMN plan_id")
    conn.execute("ALTER TABLE paradas DROP COLUMN cuadrilla")
    conn.execute("ALTER TABLE paradas DROP COLUMN fecha_informada")
    conn.execute("PRAGMA user_version = 3")
    conn.commit()
    conn.close()
    database._initialized.clear()

    estados = _estados(client)

    assert estados[("C-1", primero["id"])] == "reprogramada"
    assert estados[("C-2", primero["id"])] == "pendiente"
    assert estados[("C-1", segundo["id"])] == "pendiente"
    conn = database.connect()
    try:
        assert conn.execute("PRAGMA user_version").fetchone()[0] == database.SCHEMA_VERSION
        assert conn.execute("SELECT COUNT(*) FROM paradas WHERE plan_id IS NULL").fetchone()[0] == 0
    finally:
        conn.close()


def test_no_se_carga_reprogramada_a_mano(client: TestClient) -> None:
    """Reprogramada se deriva: no es un estado que se pueda informar."""
    _guardar(client)
    tarea = next(iter(_tareas(client).values()))

    response = client.patch(f"/registro/tareas/{tarea['id']}", json={"status": "reprogramada"})

    assert response.status_code == 422
    assert client.patch("/registro/tareas/999", json={"status": "realizada"}).status_code == 404


def test_excel_de_pendientes_entra_al_planificador(client: TestClient) -> None:
    """Lo que falta, exportado, se reimporta tal cual por el paso 1."""
    guardado = _guardar(client)
    tareas = _tareas(client)
    client.patch(f"/registro/tareas/{tareas[('C-1', guardado['id'])]['id']}", json={"status": "realizada"})
    client.patch(
        f"/registro/tareas/{tareas[('C-2', guardado['id'])]['id']}",
        json={"status": "no_realizada", "observation": "Sin acceso"},
    )

    response = client.get("/registro/tareas.xlsx", params={"estado": "faltan"})

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MIME
    assert "tareas-registro-" in response.headers["content-disposition"]
    frame = pd.read_excel(io.BytesIO(response.content))
    mapeo = suggest_mapping(list(frame.columns))
    pendientes, hechas = split_done(frame, mapeo["done"])
    puntos, descartadas = extract_points(
        pendientes, mapeo["id"], mapeo["lat"], mapeo["lon"],
        col_node=mapeo["node"], col_obs=mapeo["observation"],
    )
    # "No" no aparta la fila: lo no realizado se vuelve a planificar.
    assert hechas == 0
    assert descartadas == []
    assert sorted(puntos["id"]) == ["C-2", "C-3"]
    assert puntos.set_index("id").loc["C-2", "observation"] == "Sin acceso"


def test_excel_de_todas_las_tareas_marca_lo_realizado(client: TestClient) -> None:
    """Subido al planificador, lo realizado se aparta solo."""
    guardado = _guardar(client)
    tarea = _tareas(client)[("C-1", guardado["id"])]
    client.patch(f"/registro/tareas/{tarea['id']}", json={"status": "realizada"})

    response = client.get("/registro/tareas.xlsx")

    frame = pd.read_excel(io.BytesIO(response.content))
    assert list(frame.columns[-5:]) == ["Estado", "Plan", "Fecha de ejecución", "Cuadrilla", "Pendientes"]
    fila = frame.set_index("ID de la cámara").loc["C-1"]
    assert (fila["Realizado"], fila["Estado"]) == ("Sí", "Realizada")
    _, hechas = split_done(frame, suggest_mapping(list(frame.columns))["done"])
    assert hechas == 1


def test_excel_de_tareas_sin_resultados_da_404(client: TestClient) -> None:
    response = client.get("/registro/tareas.xlsx", params={"estado": "faltan"})

    assert response.status_code == 404


# --------------------------------------------------------------------------
# Respaldo
# --------------------------------------------------------------------------


def test_el_respaldo_es_una_base_sqlite_con_los_datos(client: TestClient) -> None:
    _guardar(client)

    response = client.get("/registro/respaldo")

    assert response.status_code == 200
    assert response.content.startswith(b"SQLite format 3\x00")
    assert ".sqlite" in response.headers["content-disposition"]
    copia = sqlite3.connect(":memory:")
    try:
        copia.deserialize(response.content)
        assert copia.execute("SELECT COUNT(*) FROM paradas").fetchone()[0] == 3
    finally:
        copia.close()


# --------------------------------------------------------------------------
# Lectura de celdas del seguimiento
# --------------------------------------------------------------------------


@pytest.mark.parametrize("valor", ["Sí", " si ", "x", "✓", 1, 1.0, True, "OK"])
def test_marcas_de_realizada(valor: object) -> None:
    assert reported_status(valor) == "realizada"


@pytest.mark.parametrize("valor", ["No", " no ", "N", 0, False, "No realizada"])
def test_marcas_de_no_realizada(valor: object) -> None:
    assert reported_status(valor) == "no_realizada"


@pytest.mark.parametrize("valor", [None, float("nan"), "", "  ", "quizás", 2])
def test_celdas_sin_novedad(valor: object) -> None:
    assert reported_status(valor) is None


def test_reconoce_las_columnas_del_excel_exportado() -> None:
    from app.services.export import FOLLOW_UP_COLUMNS

    mapeo = follow_up_mapping(FOLLOW_UP_COLUMNS)

    assert mapeo == {
        "id": "ID de la cámara",
        "done": "Realizado",
        "observation": "Observación",
        "reason": None,
        "migrated_node": "Nodo al cual se migró",
        "date": "Fecha de planificación",
        "crew": None,
        "worked_date": None,
        "truck": None,
    }


def test_la_base_de_los_tests_no_es_la_de_data() -> None:
    """La fixture autouse aísla cada test en una base temporal."""
    assert "data" not in config.DB_PATH.split("/")[-2:]


# --------------------------------------------------------------------------
# Planilla de órdenes de trabajo (indicador "Cierre")
# --------------------------------------------------------------------------

_ORDEN_GENERICA = "MIGRAR DESDE EL MARTEAREANA A CUALQUIER NODO DISPONIBLE."


def _ordenes(filas: list[dict]) -> bytes:
    """Excel de órdenes de trabajo con las columnas que importan, como lo baja el sistema."""
    base = {
        "N° OT": 0, "Estado": "Cerrada", "Fecha Programación": "5/10/2026",
        "Fecha Creación": "3/10/2026 12:35", "Observaciones": _ORDEN_GENERICA,
        "Obs. Cierre": None, "Estado No Realizado": None, "NODO": None, "COORDENADAS": None,
    }
    salida = io.BytesIO()
    pd.DataFrame(
        [{**base, "N° OT": 120000003370 + i, **fila} for i, fila in enumerate(filas)]
    ).to_excel(salida, index=False, engine="openpyxl")
    return salida.getvalue()


def _plan_de_ordenes(client: TestClient) -> dict:
    plan = _plan(
        {
            **_jornada("2026-10-05", 1, []),
            "stops": [
                _parada("A-1", node="Solidaridad"),
                _parada("A-2", node="PTP Norte"),
                _parada("A-3", node="Solidaridad"),
                _parada("A-4"),
            ],
        }
    )
    return _guardar(client, plan)


def test_las_ordenes_se_cruzan_por_id_contrato_y_no_por_numero_de_ot(client: TestClient) -> None:
    guardado = _plan_de_ordenes(client)
    contenido = _ordenes([
        {"ID Contrato": "A-1", "Cierre": "REALIZADO", "NODO": "SOLIDARIDAD"},
    ])

    resultado = _cargar(client, contenido)

    assert (resultado["matched"], resultado["unmatched"], resultado["done"]) == (1, 0, 1)
    assert _tareas(client)[("A-1", guardado["id"])]["status"] == "realizada"


def test_orden_realizada_con_el_nodo_preliminar_guarda_la_observacion(client: TestClient) -> None:
    """Aunque se migró al nodo planificado, lo que escribió el técnico puede
    importar (caso 1-0384: "no se retira el enlace previo, es para camión")."""
    guardado = _plan_de_ordenes(client)
    _cargar(client, _ordenes([{
        "ID Contrato": "A-1", "Cierre": "REALIZADO", "NODO": "SOLIDARIDAD",
        "Obs. Cierre": "Se migra a Solidaridad; no se retira el enlace previo, es para camión",
    }]))

    tarea = _tareas(client)[("A-1", guardado["id"])]
    assert (tarea["status"], tarea["migrated_node"]) == ("realizada", "SOLIDARIDAD")
    assert tarea["observation"] == "Se migra a Solidaridad; no se retira el enlace previo, es para camión"


def test_orden_realizada_en_otro_nodo_guarda_la_observacion(client: TestClient) -> None:
    guardado = _plan_de_ordenes(client)
    _cargar(client, _ordenes([{
        "ID Contrato": "A-2", "Cierre": "REALIZADO", "NODO": "SOLIDARIDAD",
        "Obs. Cierre": "Se migra de nodo y se coloca enlace nuevo",
    }]))

    tarea = _tareas(client)[("A-2", guardado["id"])]
    assert (tarea["status"], tarea["migrated_node"]) == ("realizada", "SOLIDARIDAD")
    assert tarea["observation"] == "Se migra de nodo y se coloca enlace nuevo"


def test_orden_realizada_sin_nodo_preliminar_guarda_la_observacion(client: TestClient) -> None:
    """Sin nodo con qué comparar no se puede decir que coincide: no se pierde lo escrito."""
    guardado = _plan_de_ordenes(client)
    _cargar(client, _ordenes([{
        "ID Contrato": "A-4", "Cierre": "REALIZADO", "NODO": "SOLIDARIDAD",
        "Obs. Cierre": "Se instala enlace",
    }]))

    assert _tareas(client)[("A-4", guardado["id"])]["observation"] == "Se instala enlace"


def test_orden_no_realizada_guarda_motivo_y_observacion_sin_nodo(client: TestClient) -> None:
    guardado = _plan_de_ordenes(client)
    resultado = _cargar(client, _ordenes([{
        "ID Contrato": "A-3", "Cierre": "NO REALIZADO", "Estado No Realizado": "SIN LINEA DE VISTA",
        "Obs. Cierre": "No se realiza migración: requiere una triangulación",
        "NODO": "SOLIDARIDAD",
    }]))

    tarea = _tareas(client)[("A-3", guardado["id"])]
    assert resultado["not_done"] == 1
    assert tarea["status"] == "no_realizada"
    assert tarea["observation"] == (
        "Sin linea de vista. No se realiza migración: requiere una triangulación"
    )
    assert tarea["migrated_node"] is None


def test_orden_no_realizada_no_repite_el_motivo_si_el_tecnico_lo_copio(client: TestClient) -> None:
    guardado = _plan_de_ordenes(client)
    _cargar(client, _ordenes([{
        "ID Contrato": "A-3", "Cierre": "NO REALIZADO",
        "Estado No Realizado": "FIN DE TURNO", "Obs. Cierre": "Fin de turno ",
    }]))

    assert _tareas(client)[("A-3", guardado["id"])]["observation"] == "Fin de turno"


def test_cargar_las_ordenes_dos_veces_deja_el_registro_igual(client: TestClient) -> None:
    _plan_de_ordenes(client)
    contenido = _ordenes([
        {"ID Contrato": "A-1", "Cierre": "REALIZADO", "NODO": "SOLIDARIDAD"},
        {"ID Contrato": "A-3", "Cierre": "NO REALIZADO", "Estado No Realizado": "FIN DE TURNO"},
    ])

    primera = _cargar(client, contenido)
    segunda = _cargar(client, contenido)

    assert primera["updated"] == 2
    assert segunda["updated"] == 0
    assert segunda["previously_loaded_at"] is not None


def test_reconoce_las_columnas_de_la_planilla_de_ordenes() -> None:
    columnas = [
        "N° OT", "Estado", "ID Contrato", "Fecha Programación", "Fecha Creación", "Cierre",
        "Observaciones", "Obs. Cierre", "Estado No Realizado", "NODO", "COORDENADAS",
        "Cuadrilla", "Fecha Cierre",
    ]

    assert follow_up_mapping(columnas) == {
        "id": "ID Contrato",
        "done": "Cierre",
        "observation": "Obs. Cierre",
        "reason": "Estado No Realizado",
        "migrated_node": "NODO",
        "date": "Fecha Programación",
        "crew": "Cuadrilla",
        "worked_date": "Fecha Cierre",
        "truck": None,
    }


def test_en_las_ordenes_reales_el_dia_trabajado_es_fin_atencion() -> None:
    """Así vienen las órdenes de Nubicom: sin "Fecha Cierre", pero con el
    inicio y el fin de la atención. El día trabajado es el del fin."""
    columnas = [
        "N° OT", "Estado", "ID Contrato", "Cuadrilla", "Fecha Programación", "Hora Programación",
        "Cierre", "Observaciones", "Obs. Cierre", "Fecha Creación", "Inicio Atención",
        "Fin Atención", "Estado No Realizado", "INDICAR SI REQUIERE CAMION", "NODO",
        "Técnico 1", "Técnico 2",
    ]

    mapeo = follow_up_mapping(columnas)

    assert (mapeo["date"], mapeo["worked_date"], mapeo["crew"], mapeo["truck"]) == (
        "Fecha Programación", "Fin Atención", "Cuadrilla", "INDICAR SI REQUIERE CAMION",
    )


def test_sin_fin_de_atencion_la_ejecucion_sale_del_inicio() -> None:
    columnas = ["ID Contrato", "Fecha Programación", "Cierre", "Inicio Atención"]

    assert follow_up_mapping(columnas)["worked_date"] == "Inicio Atención"


@pytest.mark.parametrize("encabezado", ["Cuadrilla", "Técnico", "Tecnico Asignado", "Móvil", "Equipo", "Brigada"])
def test_reconoce_la_columna_de_cuadrilla_con_otros_nombres(encabezado: str) -> None:
    columnas = ["ID Contrato", "Fecha Programación", "Cierre", encabezado]

    assert follow_up_mapping(columnas)["crew"] == encabezado


# --------------------------------------------------------------------------
# Cambios hechos fuera del programa: otra cuadrilla, otro día
# --------------------------------------------------------------------------


def test_una_tarea_que_otra_cuadrilla_hizo_otro_dia_queda_marcada(client: TestClient) -> None:
    """El plan era para una cuadrilla el 5/10; se liberó otra y se le pasó A-2
    para el 7/10. La carga lo cruza igual, pero lo marca y lo informa."""
    guardado = _plan_de_ordenes(client)

    resultado = _cargar(client, _ordenes([
        {"ID Contrato": "A-1", "Cierre": "REALIZADO", "Cuadrilla": "Cuadrilla 1"},
        {"ID Contrato": "A-2", "Cierre": "REALIZADO", "Cuadrilla": "Cuadrilla 2",
         "Fecha Programación": "7/10/2026"},
    ]))

    assert resultado["matched"] == 2 and resultado["done"] == 2
    assert resultado["crew_column"] == "Cuadrilla"
    assert resultado["off_plan"] == 1
    assert resultado["off_plan_tasks"] == [{
        "camera_id": "A-2", "planned_date": "2026-10-05", "reported_date": "2026-10-07",
        "crew": "Cuadrilla 2",
    }]
    tareas = _tareas(client)
    a1, a2 = tareas[("A-1", guardado["id"])], tareas[("A-2", guardado["id"])]
    assert (a1["crew"], a1["reported_date"], a1["off_plan"]) == ("Cuadrilla 1", "2026-10-05", False)
    assert (a2["crew"], a2["reported_date"], a2["off_plan"]) == ("Cuadrilla 2", "2026-10-07", True)
    assert a2["status"] == "realizada"


def test_informa_que_cuadrillas_trabajaron_cada_dia(client: TestClient) -> None:
    _plan_de_ordenes(client)

    resultado = _cargar(client, _ordenes([
        {"ID Contrato": "A-1", "Cierre": "REALIZADO", "Cuadrilla": "Cuadrilla 1"},
        {"ID Contrato": "A-2", "Cierre": "REALIZADO", "Cuadrilla": "Cuadrilla 1"},
        {"ID Contrato": "A-3", "Cierre": "NO REALIZADO", "Cuadrilla": "Cuadrilla 2"},
    ]))

    assert resultado["crews_by_day"] == [{
        "date": "2026-10-05",
        "crews": [{"crew": "Cuadrilla 1", "tasks": 2}, {"crew": "Cuadrilla 2", "tasks": 1}],
    }]
    assert resultado["off_plan"] == 0


def test_la_fecha_de_cierre_manda_sobre_la_programada(client: TestClient) -> None:
    """Si la planilla dice cuándo se cerró la orden, ése es el día en que se trabajó."""
    guardado = _plan_de_ordenes(client)

    resultado = _cargar(client, _ordenes([
        {"ID Contrato": "A-4", "Cierre": "REALIZADO", "Cuadrilla": "Cuadrilla 2", "Fecha Cierre": "6/10/2026 15:20"},
    ]))

    assert resultado["off_plan"] == 1
    assert _tareas(client)[("A-4", guardado["id"])]["reported_date"] == "2026-10-06"


def test_se_pueden_listar_solo_las_tareas_fuera_del_plan(client: TestClient) -> None:
    guardado = _plan_de_ordenes(client)
    _cargar(client, _ordenes([
        {"ID Contrato": "A-1", "Cierre": "REALIZADO", "Cuadrilla": "Cuadrilla 1"},
        {"ID Contrato": "A-2", "Cierre": "REALIZADO", "Cuadrilla": "Cuadrilla 2", "Fecha Programación": "7/10/2026"},
    ]))

    assert set(_tareas(client, estado="fuera_de_plan")) == {("A-2", guardado["id"])}


def test_una_orden_abierta_no_tiene_fecha_de_ejecucion(client: TestClient) -> None:
    """Sin cerrar, la cuadrilla todavía no pasó: aunque Zeta la haya
    reprogramado, el plan conserva su fecha y no hay ejecución que marcar."""
    guardado = _plan_de_ordenes(client)

    resultado = _cargar(client, _ordenes([
        {"ID Contrato": "A-3", "Cierre": None, "Cuadrilla": "Cuadrilla 2", "Fecha Programación": "8/10/2026"},
        {"ID Contrato": "A-4", "Cierre": None, "Cuadrilla": "Cuadrilla 1"},
    ]))

    tareas = _tareas(client)
    a3, a4 = tareas[("A-3", guardado["id"])], tareas[("A-4", guardado["id"])]
    assert resultado["off_plan"] == 0 and resultado["updated"] == 0
    assert (a3["status"], a3["crew"], a3["reported_date"], a3["off_plan"]) == ("pendiente", None, None, False)
    assert (a4["crew"], a4["verified_at"]) == (None, None)


def test_un_archivo_viejo_cargado_despues_no_pisa_lo_mas_nuevo(client: TestClient) -> None:
    """Caso 1-1494: el 05/10 no se pudo, el 06/10 la hizo la otra cuadrilla.
    Volver a subir la planilla del 05/10 no puede devolverla a "no realizada"."""
    guardado = _plan_de_ordenes(client)
    nuevo = _ordenes([{
        "ID Contrato": "A-1", "Cierre": "REALIZADO", "Cuadrilla": "AE414UC", "NODO": "SOLIDARIDAD",
        "Fecha Programación": "6/10/2026", "Fin Atención": "6/10/2026 16:04",
    }])
    viejo = _ordenes([{
        "ID Contrato": "A-1", "Cierre": "NO REALIZADO", "Cuadrilla": "AE414WY",
        "Estado No Realizado": "FIN DE TURNO", "Fin Atención": "5/10/2026 18:30",
    }])

    _cargar(client, nuevo)
    resultado = _cargar(client, viejo)

    tarea = _tareas(client)[("A-1", guardado["id"])]
    assert (tarea["status"], tarea["crew"], tarea["reported_date"]) == ("realizada", "AE414UC", "2026-10-06")
    assert resultado["stale"] == 1 and resultado["updated"] == 0


def test_sin_columna_de_cuadrilla_igual_detecta_el_cambio_de_dia(client: TestClient) -> None:
    _plan_de_ordenes(client)

    resultado = _cargar(client, _ordenes([
        {"ID Contrato": "A-2", "Cierre": "REALIZADO", "Fecha Programación": "7/10/2026"},
    ]))

    assert resultado["crew_column"] is None
    assert resultado["off_plan"] == 1
    assert resultado["off_plan_tasks"][0]["crew"] is None
    assert resultado["crews_by_day"] == []


def test_el_excel_del_registro_dice_cuando_se_ejecuto_quien_y_que_falta(client: TestClient) -> None:
    """El plan conserva sus fechas; al lado van la fecha de ejecución, la
    cuadrilla y lo que queda pendiente según la observación."""
    guardado = _plan_de_ordenes(client)
    _cargar(client, _ordenes([
        {"ID Contrato": "A-1", "Cierre": "REALIZADO", "Cuadrilla": "AE414WY", "NODO": "SOLIDARIDAD",
         "Fin Atención": "5/10/2026 16:13",
         "Obs. Cierre": "Se migra con éxito, no se retira enlace previamente instalado ya que es para camion"},
        {"ID Contrato": "A-2", "Cierre": "NO REALIZADO", "Cuadrilla": "AE414UC",
         "Estado No Realizado": "FIN DE TURNO", "Inicio Atención": "6/10/2026 17:50"},
        {"ID Contrato": "A-3", "Cierre": "REALIZADO", "Cuadrilla": "AE414WY", "NODO": "SOLIDARIDAD",
         "Fin Atención": "5/10/2026 12:00", "INDICAR SI REQUIERE CAMION": "ENLACE"},
    ]))

    tareas = _tareas(client)
    a1, a2 = tareas[("A-1", guardado["id"])], tareas[("A-2", guardado["id"])]
    assert (a1["date"], a1["reported_date"], a1["crew"]) == ("2026-10-05", "2026-10-05", "AE414WY")
    assert a1["pending_actions"] == "Retirar el enlace anterior con camión"
    assert (a2["date"], a2["reported_date"], a2["off_plan"]) == ("2026-10-05", "2026-10-06", True)
    assert a2["pending_actions"] == "Reprogramar (fin de turno)"
    assert tareas[("A-3", guardado["id"])]["pending_actions"] == "Requiere camión (enlace)"
    assert tareas[("A-4", guardado["id"])]["pending_actions"] is None

    contenido = client.get("/registro/tareas.xlsx", params={"plan_id": guardado["id"]}).content
    frame = pd.read_excel(io.BytesIO(contenido)).set_index("ID de la cámara")
    fila = frame.loc["A-2"]
    assert fila["Fecha de planificación"] == pd.Timestamp("2026-10-05")
    assert fila["Fecha de ejecución"] == pd.Timestamp("2026-10-06")
    assert (fila["Cuadrilla"], fila["Pendientes"]) == ("AE414UC", "Reprogramar (fin de turno)")
    assert pd.isna(frame.loc["A-4", "Fecha de ejecución"]) and pd.isna(frame.loc["A-4", "Cuadrilla"])

    # El mismo Excel, vuelto a subir como seguimiento, deja todo igual.
    assert _cargar(client, contenido, "registro.xlsx")["updated"] == 0


# --------------------------------------------------------------------------
# Editor completo de una tarea
# --------------------------------------------------------------------------


def _editar(client: TestClient, tarea: dict, **cambios: object):
    return client.patch(f"/registro/tareas/{tarea['id']}", json=cambios)


def test_se_editan_todos_los_datos_de_una_tarea(client: TestClient) -> None:
    guardado = _guardar(client)
    tarea = _tareas(client)[("C-1", guardado["id"])]

    response = _editar(
        client, tarea, camera_id=" C-1b ", label="Av. Belgrano 100", node="Atocha",
        lat=-24.8, lon=-65.4, status="no_realizada", reported_date="2026-10-02",
        crew="AE414UC", truck="ENLACE", migrated_node="Solidaridad", observation="Fin de turno",
    )

    assert response.status_code == 200, response.text
    editada = response.json()
    assert (editada["camera_id"], editada["label"], editada["node"]) == ("C-1b", "Av. Belgrano 100", "Atocha")
    assert (editada["lat"], editada["lon"]) == (-24.8, -65.4)
    assert (editada["status"], editada["reported_date"], editada["crew"]) == ("no_realizada", "2026-10-02", "AE414UC")
    assert editada["off_plan"] is True  # planificada el 01/10
    assert editada["pending_actions"] == "Reprogramar (fin de turno); Requiere camión (enlace)"
    campos = {fila["campo"] for fila in _correcciones()}
    assert campos == {
        "camara_id", "descripcion", "nodo_preliminar", "lat", "lon", "estado", "fecha_informada",
        "cuadrilla", "requiere_camion", "nodo_migrado", "observacion",
    }


def test_una_tarea_se_puede_pasar_a_otra_jornada_del_plan(client: TestClient) -> None:
    """Cambia la fecha planificada: va al final de la otra jornada del mismo plan."""
    guardado = _guardar(client)
    tarea = _tareas(client)[("C-1", guardado["id"])]
    otra = next(r for r in client.get("/registro/recorridos/").json() if r["date"] == "2026-10-02")

    movida = _editar(client, tarea, route_id=otra["id"]).json()

    assert (movida["route_id"], movida["date"], movida["order"]) == (otra["id"], "2026-10-02", 2)
    (fila,) = _correcciones()
    assert (fila["campo"], fila["valor_anterior"], fila["valor_nuevo"]) == (
        "jornada", "2026-10-01 (día 1)", "2026-10-02 (día 2)",
    )


def test_no_se_pasa_una_tarea_a_una_jornada_de_otro_plan(client: TestClient) -> None:
    primero = _guardar(client)
    _guardar(client, _plan(_jornada("2026-11-01", 1, ["X-1"])))
    tarea = _tareas(client)[("C-1", primero["id"])]
    ajena = next(r for r in client.get("/registro/recorridos/").json() if r["date"] == "2026-11-01")

    response = _editar(client, tarea, route_id=ajena["id"])

    assert response.status_code == 422
    assert "otro plan" in response.json()["detail"]


@pytest.mark.parametrize(
    "cambios", [{"camera_id": "   "}, {"lat": 95}, {"lon": -200}, {"reported_date": "no-es-fecha"}]
)
def test_el_editor_rechaza_datos_invalidos(client: TestClient, cambios: dict) -> None:
    guardado = _guardar(client)
    tarea = _tareas(client)[("C-1", guardado["id"])]

    assert _editar(client, tarea, **cambios).status_code == 422
    assert _correcciones() == []


def test_los_pendientes_escritos_a_mano_quedan_hasta_volver_a_automatico(client: TestClient) -> None:
    guardado = _plan_de_ordenes(client)
    tarea = _tareas(client)[("A-3", guardado["id"])]
    _cargar(client, _ordenes([{"ID Contrato": "A-3", "Cierre": "NO REALIZADO", "Estado No Realizado": "FIN DE TURNO"}]))

    a_mano = _editar(client, tarea, pending_actions="Volver con escalera larga").json()
    _cargar(client, _ordenes([{"ID Contrato": "A-3", "Cierre": "NO REALIZADO", "Estado No Realizado": "SIN LINEA DE VISTA"}]))
    tras_cargar = _tareas(client)[("A-3", guardado["id"])]
    automatico = _editar(client, tarea, pending_actions=None).json()

    assert (a_mano["pending_actions"], a_mano["pending_manual"]) == ("Volver con escalera larga", True)
    assert tras_cargar["pending_actions"] == "Volver con escalera larga"
    assert automatico["pending_manual"] is False
    assert automatico["pending_actions"].startswith("Reprogramar")


def test_una_base_v6_conserva_sus_correcciones_al_migrar(client: TestClient) -> None:
    """La tabla vieja sólo aceptaba tres campos (un CHECK que SQLite no altera):
    se rehace sin perder filas, y después acepta cualquier campo."""
    guardado = _guardar(client)
    tarea = _tareas(client)[("C-1", guardado["id"])]
    conn = database.connect()
    conn.execute("DROP TABLE correcciones")
    conn.execute(
        """CREATE TABLE correcciones (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            parada_id INTEGER REFERENCES paradas (id) ON DELETE SET NULL,
            plan_id INTEGER NOT NULL, camara_id TEXT NOT NULL, fecha TEXT NOT NULL,
            corregido_en TEXT NOT NULL, usuario TEXT NOT NULL, nombre TEXT NOT NULL,
            campo TEXT NOT NULL CHECK (campo IN ('estado', 'observacion', 'nodo_migrado')),
            valor_anterior TEXT, valor_nuevo TEXT)"""
    )
    conn.execute(
        "INSERT INTO correcciones (parada_id, plan_id, camara_id, fecha, corregido_en, usuario, "
        "nombre, campo, valor_anterior, valor_nuevo) VALUES (?, ?, 'C-1', '2026-10-01', 'x', "
        "'admin', 'Admin', 'estado', 'pendiente', 'realizada')",
        (tarea["id"], guardado["id"]),
    )
    conn.execute("ALTER TABLE paradas DROP COLUMN pendientes_manual")
    conn.execute("PRAGMA user_version = 6")
    conn.commit()
    conn.close()
    database._initialized.clear()

    assert _editar(client, tarea, lat=-24.81).status_code == 200

    assert [(f["usuario"], f["campo"]) for f in _correcciones()] == [("admin", "estado"), ("pruebas", "lat")]


# --------------------------------------------------------------------------
# Editar una jornada entera
# --------------------------------------------------------------------------


def _jornada_del(client: TestClient, fecha: str) -> dict:
    return next(r for r in client.get("/registro/recorridos/").json() if r["date"] == fecha)


def test_se_cambia_la_fecha_planificada_de_un_dia_entero(client: TestClient) -> None:
    guardado = _guardar(client)
    jornada = _jornada_del(client, "2026-10-01")

    response = client.patch(f"/registro/recorridos/{jornada['id']}", json={"date": "2026-10-09"})

    assert response.status_code == 200, response.text
    assert response.json()["date"] == "2026-10-09"
    tareas = _tareas(client)
    assert {tareas[(c, guardado["id"])]["date"] for c in ("C-1", "C-2")} == {"2026-10-09"}
    assert tareas[("C-3", guardado["id"])]["date"] == "2026-10-02"  # la otra jornada no cambia
    filas = _correcciones()
    assert {(f["camara_id"], f["campo"], f["valor_anterior"], f["valor_nuevo"]) for f in filas} == {
        ("C-1", "fecha_planificada", "2026-10-01", "2026-10-09"),
        ("C-2", "fecha_planificada", "2026-10-01", "2026-10-09"),
    }


def test_cambiar_el_dia_recalcula_lo_que_quedo_fuera_del_plan(client: TestClient) -> None:
    """Si el día se movió a cuando de verdad se trabajó, deja de estar fuera del plan."""
    guardado = _guardar(client)
    tarea = _tareas(client)[("C-1", guardado["id"])]
    _editar(client, tarea, status="realizada", reported_date="2026-10-03")
    assert _tareas(client)[("C-1", guardado["id"])]["off_plan"] is True

    client.patch(f"/registro/recorridos/{tarea['route_id']}", json={"date": "2026-10-03"})

    assert _tareas(client)[("C-1", guardado["id"])]["off_plan"] is False


def test_mover_un_dia_es_solo_para_administradores(client: TestClient, como_operador) -> None:
    _guardar(client)
    jornada = _jornada_del(client, "2026-10-01")

    assert client.patch(f"/registro/recorridos/{jornada['id']}", json={"date": "2026-10-09"}).status_code == 403


def test_mover_un_dia_valida_la_fecha_y_la_jornada(client: TestClient) -> None:
    _guardar(client)
    jornada = _jornada_del(client, "2026-10-01")

    assert client.patch(f"/registro/recorridos/{jornada['id']}", json={"date": "09/10"}).status_code == 422
    assert client.patch("/registro/recorridos/999", json={"date": "2026-10-09"}).status_code == 404
    assert _correcciones() == []
