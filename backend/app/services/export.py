"""Excel de seguimiento del plan: lo completan los técnicos y vuelve a la app.

La primera hoja ("Recorridos") tiene una fila por cámara con las columnas que
llenan los técnicos. Es también la que lee la ingesta al reimportar el archivo:
por eso los encabezados van en la fila 1, sin títulos arriba, y se incluyen
latitud y longitud — sin ellas las tareas pendientes no se podrían volver a
planificar sin la planilla original. La segunda hoja ("Resumen") agrupa por
día para estimar tiempos. Si el plan está guardado en el registro, una hoja
oculta ("_registro") lleva su id: al cargar el seguimiento se actualiza ese
plan y no otro con la misma cámara y fecha.

El mismo formato sirve para exportar tareas del registro (`build_tasks_workbook`):
lo pendiente vuelve a entrar por el paso 1 del planificador tal cual.
"""

import datetime as dt
import io

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from ..schemas import ExportDay, ExportRequest, Task

FOLLOW_UP_COLUMNS = [
    "ID de la cámara",
    "Fecha de planificación",
    "Orden",
    "Realizado",
    "Observación",
    "Nodo al cual se migró",
    "Nodo preliminar",
    "Latitud",
    "Longitud",
]
_WIDTHS = [18, 14, 8, 12, 40, 26, 26, 12, 12]
_REPLAN_COLUMNS = {"Latitud", "Longitud"}
_DONE_OPTIONS = '"Sí,No"'

# Columnas extra del Excel de tareas del registro. Van al final para no mover
# las que reconoce `suggest_mapping` al reimportarlo.
TASK_EXTRA_COLUMNS = ["Estado", "Plan"]
_TASK_EXTRA_WIDTHS = [16, 30]
_STATUS_TEXT = {
    "pendiente": "Pendiente",
    "realizada": "Realizada",
    "no_realizada": "No realizada",
    "reprogramada": "Reprogramada",
}
# Lo que va en "Realizado" según el estado: lo pendiente queda vacío para que
# el técnico lo complete y para que el planificador lo vuelva a tomar.
_DONE_TEXT = {"realizada": "Sí", "no_realizada": "No"}

# Hoja oculta con el id del plan del registro.
REGISTRY_SHEET = "_registro"

SUMMARY_COLUMNS = [
    "Fecha",
    "Día",
    "Cluster",
    "Salida",
    "Cámaras",
    "Distancia (km)",
    "Duración estimada",
]
_SUMMARY_WIDTHS = [14, 8, 10, 24, 10, 16, 18]

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1E293B")
_REPLAN_FILL = PatternFill("solid", fgColor="64748B")
# Banda suave en días alternos: separa visualmente un día del siguiente.
_BAND_FILL = PatternFill("solid", fgColor="F1F5F9")
_CENTER = Alignment(horizontal="center", vertical="center")
_DATE_FORMAT = "DD/MM/YYYY"


def _set_text(cell: Cell, text: str | None) -> None:
    """Escribe texto literal. openpyxl toma como fórmula todo lo que empieza
    con "=": un ID u observación así se ejecutaría al abrir el archivo."""
    cell.value = text
    if isinstance(text, str):
        cell.data_type = "s"


def _style_header(sheet: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = _HEADER_FONT
        cell.fill = _REPLAN_FILL if cell.value in _REPLAN_COLUMNS else _HEADER_FILL
        cell.alignment = _CENTER
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.row_dimensions[1].height = 22
    sheet.freeze_panes = "A2"


def _sorted_days(plan: ExportRequest) -> list[ExportDay]:
    return sorted(plan.days, key=lambda day: (day.date, day.cluster_id, day.day))


def _append_follow_up_row(
    sheet: Worksheet,
    *,
    camera_id: str,
    date: dt.date,
    order: int,
    lat: float,
    lon: float,
    node: str | None,
    observation: str | None,
    done: str | None = None,
    migrated_node: str | None = None,
    extra: list[str | None] | None = None,
    banded: bool = False,
) -> None:
    """Agrega una fila con las columnas de FOLLOW_UP_COLUMNS (y `extra` al final)."""
    row = sheet.max_row + 1
    _set_text(sheet.cell(row=row, column=1), camera_id)
    sheet.cell(row=row, column=2, value=date).number_format = _DATE_FORMAT
    sheet.cell(row=row, column=3, value=order).alignment = _CENTER
    _set_text(sheet.cell(row=row, column=4), done)
    sheet.cell(row=row, column=4).alignment = _CENTER
    _set_text(sheet.cell(row=row, column=5), observation)
    _set_text(sheet.cell(row=row, column=6), migrated_node)
    _set_text(sheet.cell(row=row, column=7), node)
    sheet.cell(row=row, column=8, value=lat)
    sheet.cell(row=row, column=9, value=lon)
    for offset, text in enumerate(extra or [], start=len(FOLLOW_UP_COLUMNS) + 1):
        _set_text(sheet.cell(row=row, column=offset), text)
    if banded:
        for column in range(1, sheet.max_column + 1):
            sheet.cell(row=row, column=column).fill = _BAND_FILL


def _fill_follow_up(sheet: Worksheet, days: list[ExportDay]) -> None:
    sheet.title = "Recorridos"
    sheet.append(FOLLOW_UP_COLUMNS)
    for band, day in enumerate(days):
        for order, stop in enumerate(day.stops, start=1):
            _append_follow_up_row(
                sheet, camera_id=stop.camera_id, date=day.date, order=order,
                lat=stop.lat, lon=stop.lon, node=stop.node,
                observation=stop.observation, banded=bool(band % 2),
            )
    _finish_follow_up(sheet, _WIDTHS)


def _finish_follow_up(sheet: Worksheet, widths: list[int]) -> None:
    """Encabezado, filtro, desplegable de Realizado y aviso en Latitud/Longitud."""
    _style_header(sheet, widths)
    last_row = sheet.max_row
    sheet.auto_filter.ref = f"A1:{sheet.cell(row=1, column=len(widths)).column_letter}{last_row}"

    done = DataValidation(
        type="list", formula1=_DONE_OPTIONS, allow_blank=True,
        showErrorMessage=True, errorTitle="Realizado",
        error="Elegí Sí o No (o dejalo vacío si no se hizo).",
    )
    done.add(f"D2:D{last_row}")
    sheet.add_data_validation(done)

    for column in (8, 9):
        sheet.cell(row=1, column=column).comment = Comment(
            "Necesaria para volver a planificar las tareas pendientes. "
            "No la modifiques.",
            "Sistema de recorridos",
        )


def _fill_summary(sheet: Worksheet, days: list[ExportDay]) -> None:
    sheet.append(SUMMARY_COLUMNS)
    for day in days:
        row = sheet.max_row + 1
        sheet.cell(row=row, column=1, value=day.date).number_format = _DATE_FORMAT
        sheet.cell(row=row, column=2, value=day.day).alignment = _CENTER
        sheet.cell(row=row, column=3, value=day.cluster_id).alignment = _CENTER
        _set_text(sheet.cell(row=row, column=4), day.start_name)
        sheet.cell(row=row, column=5, value=len(day.stops)).alignment = _CENTER
        sheet.cell(row=row, column=6, value=round(day.distance_m / 1000, 1)).number_format = "0.0"
        # Fracción de día con formato [h]:mm: Excel la suma como duración.
        duration = sheet.cell(row=row, column=7, value=day.duration_s / 86400)
        duration.number_format = "[h]:mm"

    total = sheet.max_row + 1
    sheet.cell(row=total, column=1, value="Total").font = Font(bold=True)
    for column, letter, number_format in ((5, "E", "0"), (6, "F", "0.0"), (7, "G", "[h]:mm")):
        cell = sheet.cell(row=total, column=column, value=f"=SUM({letter}2:{letter}{total - 1})")
        cell.number_format = number_format
        cell.font = Font(bold=True)
    _style_header(sheet, _SUMMARY_WIDTHS)


def _save(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_plan_workbook(plan: ExportRequest) -> bytes:
    """Arma el .xlsx de seguimiento y lo devuelve como bytes."""
    days = _sorted_days(plan)
    workbook = Workbook()
    _fill_follow_up(workbook.active, days)
    _fill_summary(workbook.create_sheet("Resumen"), days)
    if plan.plan_id is not None:
        registry = workbook.create_sheet(REGISTRY_SHEET)
        registry.append(["plan_id", plan.plan_id])
        registry.sheet_state = "hidden"
    return _save(workbook)


def read_plan_id(filename: str, raw: bytes) -> int | None:
    """Id del plan del registro guardado en la hoja oculta, si el archivo la tiene.

    Un CSV, un Excel ajeno o uno que alguien rehízo desde cero no la tienen: en
    ese caso el seguimiento se cruza sólo por cámara y fecha.
    """
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        return None
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True)
    except Exception:  # openpyxl levanta tipos muy variados ante un archivo roto
        return None
    try:
        if REGISTRY_SHEET not in workbook.sheetnames:
            return None
        for key, value, *_ in workbook[REGISTRY_SHEET].iter_rows(max_row=5, values_only=True):
            if key == "plan_id" and isinstance(value, (int, float)) and value >= 1:
                return int(value)
        return None
    finally:
        workbook.close()


def build_tasks_workbook(tasks: list[Task]) -> bytes:
    """Excel de tareas del registro, en el formato del seguimiento.

    "Realizado" lleva Sí/No según lo informado y queda vacío en lo pendiente:
    subido al planificador, lo hecho se aparta y el resto se vuelve a planificar.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tareas"
    sheet.append(FOLLOW_UP_COLUMNS + TASK_EXTRA_COLUMNS)
    ordered = sorted(tasks, key=lambda t: (t.date, t.plan_id, t.cluster_id, t.day, t.order))
    band = -1
    previous_route: int | None = None
    for task in ordered:
        if task.route_id != previous_route:
            band += 1
            previous_route = task.route_id
        _append_follow_up_row(
            sheet, camera_id=task.camera_id, date=task.date, order=task.order,
            lat=task.lat, lon=task.lon, node=task.node, observation=task.observation,
            done=_DONE_TEXT.get(task.status), migrated_node=task.migrated_node,
            extra=[_STATUS_TEXT[task.status], task.plan_name],
            banded=bool(band % 2),
        )
    _finish_follow_up(sheet, _WIDTHS + _TASK_EXTRA_WIDTHS)
    return _save(workbook)


def tasks_filename(today: dt.date) -> str:
    return f"tareas-registro-{today.isoformat()}.xlsx"


def export_filename(plan: ExportRequest) -> str:
    """Nombre sugerido para la descarga, según la primera fecha del plan."""
    first: dt.date = min(day.date for day in plan.days)
    return f"plan-recorridos-{first.isoformat()}.xlsx"
